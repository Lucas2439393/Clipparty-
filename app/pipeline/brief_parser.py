"""
Leest een campagnebriefing (pdf/docx/txt) en extraheert er een CampaignProfile uit.
Regel: als iets niet in de briefing staat, wordt het NIET verzonnen -- het veld
blijft leeg/None. Elke nieuwe campagne is zelfstandig; er wordt niets hergebruikt
van eerdere campagnes.
"""
import json
from pathlib import Path

import pdfplumber
from docx import Document

from app.models import CampaignProfile
from app.pipeline import llm_provider

EXTRACTION_SYSTEM_PROMPT = """Je bent een strikte briefing-extractor voor een short-form \
video clipping workflow. Je krijgt de volledige tekst van een campagnebriefing.

Extraheer UITSLUITEND informatie die daadwerkelijk in de tekst staat. Verzin NOOIT \
waarden die er niet in staan -- laat het veld dan leeg (null, of lege lijst).

Geef ALLEEN geldig JSON terug, geen uitleg, geen markdown-fences, met exact deze velden:
{
  "merk": string|null,
  "campagne": string|null,
  "doel": string|null,
  "doelgroep": string|null,
  "bronmateriaal": string|null,
  "toegestane_onderwerpen": string[],
  "verboden_onderwerpen": string[],
  "hook_richtlijnen": string|null,
  "edit_richtlijnen": string|null,
  "minimumduur": number|null,
  "maximumduur": number|null,
  "captions": string|null,
  "hashtags": string[],
  "tags": string[],
  "links": string[],
  "disclaimers": string|null,
  "platformregels": string|null,
  "verboden_claims": string[],
  "overige_compliance_eisen": string|null,
  "toegestane_bronvideos": string[],
  "raw_extraction_notes": string|null
}

"maximumduur" mag nooit hoger dan 60 worden geëxtraheerd, ook al staat er iets hogers \
in de tekst -- als de briefing dat toch doet, zet "maximumduur" op 60 en leg dit uit \
in raw_extraction_notes.
"""


def _read_pdf(path: str) -> str:
    text_parts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            t = page.extract_text() or ""
            text_parts.append(t)
    return "\n".join(text_parts)


def _read_docx(path: str) -> str:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _read_txt(path: str) -> str:
    return Path(path).read_text(encoding="utf-8", errors="ignore")


def read_brief_text(path: str) -> str:
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return _read_pdf(path)
    if ext == ".docx":
        return _read_docx(path)
    if ext == ".txt":
        return _read_txt(path)
    raise ValueError(f"Niet-ondersteund briefingformaat: {ext} (verwacht .pdf, .docx of .txt)")


def extract_campaign_profile(
    brief_path: str, provider: str = "groq",
) -> tuple[CampaignProfile, str, str | None, str | None]:
    """
    Leest en extraheert de briefing met de gekozen LLM-provider ("groq" of
    "anthropic"). Retourneert (profile, provider_used, fallback_reason, model_used).
    fallback_reason is niet None als "anthropic" gevraagd was maar er
    automatisch is teruggeschakeld naar Groq omdat ANTHROPIC_API_KEY ontbreekt.
    model_used is het daadwerkelijk gebruikte Groq-model (automatisch bepaald,
    nooit hardcoded) of het Anthropic-model.
    """
    provider_used, fallback_reason = llm_provider.resolve_provider(provider)

    raw_text = read_brief_text(brief_path)
    if not raw_text.strip():
        raise ValueError(f"Kon geen tekst uit briefing halen: {brief_path}")

    text, model_used = llm_provider.complete(
        system=EXTRACTION_SYSTEM_PROMPT,
        user=raw_text[:60000],
        provider=provider_used,
        max_tokens=1800,
        json_mode=True,
    )
    text = llm_provider.strip_json_fences(text)

    try:
        data = json.loads(text)
    except json.JSONDecodeError as e:
        raise RuntimeError(f"Kon briefing-extractie niet parsen als JSON: {e}\nRuwe output: {text[:500]}")

    # Cap maximumduur hard op 60, ongeacht wat het model teruggeeft.
    if data.get("maximumduur") and data["maximumduur"] > 60:
        data["maximumduur"] = 60

    return CampaignProfile(**data), provider_used, fallback_reason, model_used
